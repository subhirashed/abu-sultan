"""Builds leads.xlsx (lightweight CRM) and leads.csv. Run: python3 build_tracker.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import csv, os
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
wb = Workbook()

# ---------- Leads ----------
ws = wb.active; ws.title = "Leads"
cols = ["Lead ID","Date received","Source","Customer name","Phone","Contact via","Service requested",
        "Pickup location (zip/city)","Drop-off location","Move date","Helpers","Photos received? (Y/N)",
        "Quoted price $","Qualified? (Y/N)","Quote sent (date)","Status","Booked date",
        "Job done (date)","Revenue $","Review asked?","Referral source","Notes"]
ws.append(cols)
hdr = PatternFill("solid", fgColor="0F6B3A")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
ws.freeze_panes = "A2"
widths = [8,13,13,18,14,11,22,18,18,12,9,9,13,11,13,13,12,12,11,11,14,30]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
# Example row (delete when real leads arrive)
ws.append(["L001",date(2026,10,17),"Craigslist","EXAMPLE – delete","305-000-0000","SMS","Unload U-Haul",
           "33020 Hollywood","same",date(2026,10,25),1,"Y",150,"Y",date(2026,10,17),"Quoted","","","","N","",""])
# Validation lists
src = DataValidation(type="list", formula1='"Craigslist,Facebook group,Facebook ad,Facebook Marketplace,Nextdoor,Google,U-Haul Moving Help,TaskRabbit,Referral,Partner,Repeat,Other"', allow_blank=True)
via = DataValidation(type="list", formula1='"SMS,WhatsApp,Call,Facebook DM,Email,Other"', allow_blank=True)
yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
status = DataValidation(type="list", formula1='"New,Replied,Quoted,Booked,Done,Lost,No response,Spam"', allow_blank=True)
for dv,col in [(src,"C"),(via,"F"),(yn,"L"),(yn,"N"),(status,"P"),(yn,"T")]:
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}1000")

# ---------- Campaign Log ----------
cl = wb.create_sheet("Campaign Log")
cl.append(["Date","Channel","Where (sub-area / group / ad set)","Title / angle","Cost $","Live until","Leads attributed","Notes"])
for c in cl[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
cl.freeze_panes = "A2"
for i,w in enumerate([12,16,34,34,8,12,14,30],1): cl.column_dimensions[get_column_letter(i)].width = w
cl.append(["","Craigslist","brw – services > labor/move","Angle 1 – Text me first",5,"","",'Status: NOT POSTED – awaiting approval'])
cl.append(["","Facebook ad","25 mi radius – WhatsApp objective","Angle 1 creative",10,"","",'Status: NOT LAUNCHED – awaiting approval'])

# ---------- Weekly Scorecard ----------
sc = wb.create_sheet("Weekly Scorecard")
sc.append(["Week start","Money spent $","Leads","Qualified leads","Quotes sent","Jobs booked","Revenue $","Cost / lead $","Cost / booked job $","Best source","Notes"])
for c in sc[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
sc.freeze_panes = "A2"
for i,w in enumerate([12,13,8,14,11,11,11,12,16,16,30],1): sc.column_dimensions[get_column_letter(i)].width = w
for r in range(2,54):
    sc[f"B{r}"] = f"=SUMIFS('Campaign Log'!E:E,'Campaign Log'!A:A,\">=\"&A{r},'Campaign Log'!A:A,\"<\"&A{r}+7)"
    sc[f"C{r}"] = f"=COUNTIFS(Leads!B:B,\">=\"&A{r},Leads!B:B,\"<\"&A{r}+7)"
    sc[f"D{r}"] = f"=COUNTIFS(Leads!B:B,\">=\"&A{r},Leads!B:B,\"<\"&A{r}+7,Leads!N:N,\"Y\")"
    sc[f"E{r}"] = f"=COUNTIFS(Leads!O:O,\">=\"&A{r},Leads!O:O,\"<\"&A{r}+7)"
    sc[f"F{r}"] = f"=COUNTIFS(Leads!Q:Q,\">=\"&A{r},Leads!Q:Q,\"<\"&A{r}+7)"
    sc[f"G{r}"] = f"=SUMIFS(Leads!S:S,Leads!R:R,\">=\"&A{r},Leads!R:R,\"<\"&A{r}+7)"
    sc[f"H{r}"] = f'=IF(C{r}>0,B{r}/C{r},"")'
    sc[f"I{r}"] = f'=IF(F{r}>0,B{r}/F{r},"")'
sc["A2"] = date(2026,10,20)
for r in range(3,54): sc[f"A{r}"] = f"=A{r-1}+7"
for r in range(2,54): sc[f"A{r}"].number_format = "yyyy-mm-dd"
for col in "BJORQ":
    for r in range(2,1001): ws[f"{col}{r}"].number_format = "yyyy-mm-dd"
for r in range(2,200): cl[f"A{r}"].number_format = "yyyy-mm-dd"; cl[f"F{r}"].number_format = "yyyy-mm-dd"

# ---------- Source Performance ----------
sp = wb.create_sheet("Source Performance")
sp.append(["Source","Leads","Qualified","Booked","Revenue $","Spend $","Cost / booked job $"])
for c in sp[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
sources = ["Craigslist","Facebook group","Facebook ad","Facebook Marketplace","Nextdoor","Google","U-Haul Moving Help","TaskRabbit","Referral","Partner","Repeat","Other"]
for i,s in enumerate(sources, start=2):
    sp[f"A{i}"] = s
    sp[f"B{i}"] = f'=COUNTIF(Leads!C:C,A{i})'
    sp[f"C{i}"] = f'=COUNTIFS(Leads!C:C,A{i},Leads!N:N,"Y")'
    sp[f"D{i}"] = f'=COUNTIFS(Leads!C:C,A{i},Leads!P:P,"Booked")+COUNTIFS(Leads!C:C,A{i},Leads!P:P,"Done")'
    sp[f"E{i}"] = f'=SUMIF(Leads!C:C,A{i},Leads!S:S)'
    sp[f"F{i}"] = f"=SUMIF('Campaign Log'!B:B,A{i},'Campaign Log'!E:E)"
    sp[f"G{i}"] = f'=IF(D{i}>0,F{i}/D{i},"")'
for i,w in enumerate([22,8,10,8,11,9,18],1): sp.column_dimensions[get_column_letter(i)].width = w

# ---------- Follow-up rules ----------
fu = wb.create_sheet("Follow-up Rules")
rules = [["When","Action","Template (07-templates/customer-reply-templates.md)"],
 ["New lead arrives","Reply within 5 minutes, ask the 5 quote questions","T1 First reply"],
 ["Info received","Send price same conversation","T2 Quote"],
 ["Quote sent, no answer 4 h","One nudge","T3 Nudge"],
 ["Quote sent, no answer 24 h","Final nudge, then mark 'No response'","T4 Final nudge"],
 ["Booked","Confirm date/time/address/payment","T5 Booking confirmation"],
 ["Day before job","Reminder","T6 Reminder"],
 ["Job done","Thank + ask for review/referral","T7 Review ask"],
 ["30 days after job","Check-in for referrals","T8 Referral check-in"]]
for r in rules: fu.append(r)
for c in fu[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
for i,w in enumerate([28,46,40],1): fu.column_dimensions[get_column_letter(i)].width = w

wb.save(os.path.join(HERE, "leads.xlsx"))
with open(os.path.join(HERE, "leads.csv"), "w", newline="") as f:
    csv.writer(f).writerow(cols)
print("built leads.xlsx and leads.csv")
